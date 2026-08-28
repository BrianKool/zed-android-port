#!/usr/bin/env python3
"""Structured, atomic Office operations for Zdroid-B agents."""

import json
import os
import sys
import tempfile
from pathlib import Path


def fail(code, message):
    print(json.dumps({"ok": False, "code": code, "message": message}))
    raise SystemExit(1)


def atomic_output(path, writer, validator):
    target = Path(path).expanduser().resolve()
    if target.exists():
        fail("output_exists", f"Refusing to overwrite existing output: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{target.name}.", suffix=target.suffix, dir=target.parent)
    os.close(fd)
    try:
        writer(temporary)
        validator(temporary)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return str(target)


def inspect_excel(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        return {"sheets": [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column} for ws in book.worksheets]}
    finally:
        book.close()


def extract_excel(path, options):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[options.get("sheet") or book.sheetnames[0]]
        limit = max(1, min(int(options.get("limit", 500)), 5000))
        rows = list(sheet.iter_rows(values_only=True))[:limit]
        return {"sheet": sheet.title, "rows": rows, "truncated": sheet.max_row > limit}
    finally:
        book.close()


def excel_frame(path, sheet_selector):
    import openpyxl
    import polars as pl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet_selector, int):
            if sheet_selector < 0 or sheet_selector >= len(book.sheetnames):
                fail("invalid_sheet", f"Sheet index is outside 0..{len(book.sheetnames) - 1}")
            sheet = book[book.sheetnames[sheet_selector]]
        else:
            sheet = book[sheet_selector or book.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            fail("empty_workbook", f"{path} contains no header row")
        headers = []
        counts = {}
        for index, raw in enumerate(raw_headers):
            base = str(raw).strip() if raw is not None and str(raw).strip() else f"column_{index + 1}"
            counts[base] = counts.get(base, 0) + 1
            headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
        values = [list(row) for row in rows]
        return pl.DataFrame(values, schema=headers, orient="row", infer_schema_length=1000)
    finally:
        book.close()


def reconcile_excel(left, right, output, options):
    import polars as pl
    import xlsxwriter

    left_sheet = options.get("left_sheet", 0)
    right_sheet = options.get("right_sheet", 0)
    left_df = excel_frame(left, left_sheet)
    right_df = excel_frame(right, right_sheet)
    keys = options.get("keys") or []
    if not keys:
        common = [name for name in left_df.columns if name in right_df.columns]
        if not common:
            fail("no_matching_columns", "No common columns were found; provide options.keys")
        keys = [common[0]]
    missing = [key for key in keys if key not in left_df.columns or key not in right_df.columns]
    if missing:
        fail("missing_keys", f"Reconciliation keys are missing: {missing}")

    left_only = left_df.join(right_df.select(keys), on=keys, how="anti")
    right_only = right_df.join(left_df.select(keys), on=keys, how="anti")
    matched = left_df.join(right_df, on=keys, how="inner", suffix="_right")
    summary = [
        ("Matched", matched.height),
        ("Only in left", left_only.height),
        ("Only in right", right_only.height),
    ]

    def write(temp):
        workbook = xlsxwriter.Workbook(
            temp,
            {"strings_to_formulas": False, "strings_to_urls": False},
        )
        header = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
        formats = {
            "Matched": workbook.add_format({"bg_color": "#E2F0D9"}),
            "Only in left": workbook.add_format({"bg_color": "#FCE4D6"}),
            "Only in right": workbook.add_format({"bg_color": "#FFF2CC"}),
        }

        def add_frame(name, frame, status):
            ws = workbook.add_worksheet(name[:31])
            columns = frame.columns + ["Reconciliation Status"]
            for col, value in enumerate(columns):
                ws.write(0, col, value, header)
            for row, values in enumerate(frame.iter_rows(), 1):
                for col, value in enumerate(values):
                    ws.write(row, col, value)
                ws.write(row, len(frame.columns), status, formats[status])
            ws.freeze_panes(1, 0)
            if columns:
                ws.autofilter(0, 0, max(frame.height, 1), len(columns) - 1)
            ws.set_column(0, max(len(columns) - 1, 0), 18)

        summary_ws = workbook.add_worksheet("Summary")
        summary_ws.write_row(0, 0, ["Status", "Count"], header)
        for row, entry in enumerate(summary, 1):
            summary_ws.write_row(row, 0, entry)
        chart = workbook.add_chart({"type": "pie"})
        chart.add_series({
            "name": "Reconciliation",
            "categories": "=Summary!$A$2:$A$4",
            "values": "=Summary!$B$2:$B$4",
            "data_labels": {"percentage": True},
        })
        chart.set_title({"name": "Reconciliation result"})
        summary_ws.insert_chart("D2", chart)
        add_frame("Matched", matched, "Matched")
        add_frame("Only in left", left_only, "Only in left")
        add_frame("Only in right", right_only, "Only in right")
        workbook.close()

    result = atomic_output(output, write, inspect_excel)
    return {"output": result, "keys": keys, "counts": dict(summary), "validation": inspect_excel(result)}


def inspect_word(path):
    from docx import Document
    doc = Document(path)
    return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "characters": sum(len(p.text) for p in doc.paragraphs)}


def extract_word(path, options):
    from docx import Document
    doc = Document(path)
    paragraph_limit = max(1, min(int(options.get("paragraph_limit", 1000)), 5000))
    table_limit = max(0, min(int(options.get("table_limit", 100)), 500))
    return {
        "paragraphs": [p.text for p in doc.paragraphs[:paragraph_limit]],
        "tables": [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables[:table_limit]],
        "truncated": len(doc.paragraphs) > paragraph_limit or len(doc.tables) > table_limit,
    }


def create_word(output, options):
    from docx import Document
    def write(temp):
        doc = Document()
        if options.get("title"):
            doc.add_heading(options["title"], 0)
        for block in options.get("blocks", []):
            kind = block.get("type", "paragraph")
            if kind == "heading": doc.add_heading(block.get("text", ""), int(block.get("level", 1)))
            elif kind == "table":
                rows = block.get("rows", [])
                if rows:
                    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                    for r, values in enumerate(rows):
                        for c, value in enumerate(values): table.cell(r, c).text = str(value)
            else: doc.add_paragraph(block.get("text", ""))
        doc.save(temp)
    result = atomic_output(output, write, inspect_word)
    return {"output": result, "validation": inspect_word(result)}


def inspect_presentation(path):
    from pptx import Presentation
    deck = Presentation(path)
    return {"slides": len(deck.slides), "shapes": sum(len(slide.shapes) for slide in deck.slides)}


def extract_presentation(path, options):
    from itertools import islice
    from pptx import Presentation
    deck = Presentation(path)
    limit = max(1, min(int(options.get("limit", 100)), 300))
    return {
        "slides": [[shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text] for slide in islice(deck.slides, limit)],
        "truncated": len(deck.slides) > limit,
    }


def create_presentation(output, options):
    from pptx import Presentation
    def write(temp):
        deck = Presentation()
        for slide_data in options.get("slides", []):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = slide_data.get("title", "")
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = slide_data.get("body", "")
        deck.save(temp)
    result = atomic_output(output, write, inspect_presentation)
    return {"output": result, "validation": inspect_presentation(result)}


def inspect_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    return {"pages": len(reader.pages), "encrypted": reader.is_encrypted}


def extract_pdf(path, options):
    from pypdf import PdfReader
    reader = PdfReader(path)
    limit = max(1, min(int(options.get("limit", 50)), len(reader.pages), 200))
    character_limit = max(1000, min(int(options.get("character_limit", 100000)), 500000))
    pages = []
    characters = 0
    truncated = len(reader.pages) > limit
    for i in range(limit):
        text = reader.pages[i].extract_text() or ""
        remaining = character_limit - characters
        if remaining <= 0:
            truncated = True
            break
        if len(text) > remaining:
            text = text[:remaining]
            truncated = True
        pages.append({"page": i + 1, "text": text})
        characters += len(text)
    return {"pages": pages, "truncated": truncated}


def create_pdf(output, options):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    def write(temp):
        c = canvas.Canvas(temp, pagesize=A4)
        width, height = A4
        y = height - 56
        for line in options.get("lines", []):
            if y < 56:
                c.showPage(); y = height - 56
            c.drawString(56, y, str(line)[:160]); y -= 18
        c.save()
    result = atomic_output(output, write, inspect_pdf)
    return {"output": result, "validation": inspect_pdf(result)}


def main():
    request = json.load(sys.stdin)
    plugin = request.get("plugin")
    action = request.get("action")
    source = request.get("input_path")
    second = request.get("second_input_path")
    output = request.get("output_path")
    options = request.get("options") or {}
    if plugin == "excel":
        result = inspect_excel(source) if action == "inspect" else extract_excel(source, options) if action == "extract" else reconcile_excel(source, second, output, options) if action == "reconcile" else fail("unsupported_action", action)
    elif plugin == "word":
        result = inspect_word(source) if action == "inspect" else extract_word(source, options) if action == "extract" else create_word(output, options) if action == "create" else fail("unsupported_action", action)
    elif plugin == "powerpoint":
        result = inspect_presentation(source) if action == "inspect" else extract_presentation(source, options) if action == "extract" else create_presentation(output, options) if action == "create" else fail("unsupported_action", action)
    elif plugin == "pdf":
        result = inspect_pdf(source) if action == "inspect" else extract_pdf(source, options) if action == "extract" else create_pdf(output, options) if action == "create" else fail("unsupported_action", action)
    else:
        fail("invalid_plugin", "plugin must be excel, word, powerpoint, or pdf")
    print(json.dumps({"ok": True, "plugin": plugin, "action": action, "result": result}, default=str))


if __name__ == "__main__":
    main()
